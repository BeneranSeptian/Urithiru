#!/usr/bin/env python3
"""Creates the tracking documentation workbook from a tracking-spec JSON file.

  pip install openpyxl
  python3 make_workbook.py <file.json>                 # writes tracking_workbook.xlsx
  python3 make_workbook.py <file.json> --out name.xlsx --blocks 25 --rows 8

Sheets:
  HOW TO USE     short instructions
  FLOW TEMPLATE  empty, formatted screen blocks; copy it for every new flow
  SUMMARY        every key from the JSON, one row each, with a Done column to tick by hand
  PROGRESS       live counts of ticked keys, overall and per feature

Everything (trackers, key prefixes, features) is read from the JSON; nothing project-specific
is written in this script.
"""
import argparse
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    sys.exit("needs openpyxl:  pip install openpyxl")

F = "Arial"
INK, MUTED, LINE = "1E2433", "5B6478", "C9CED8"
HEAD_BG, BLOCK_BG, BLOCK_INK = "1E2433", "DCE3EE", "1E2433"
THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True)
MID = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DONE_MARK = "✓"


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


# ---- reading + grouping the JSON ----------------------------------------------------------------
def split_key(k):
    m = re.match(r"^([a-z]+)(.*)$", k)
    return (m.group(1), m.group(2)) if m else ("", k)


def tokens(s):
    return re.findall(r"[A-Z]+(?=[A-Z][a-z]|\d|$)|[A-Z]?[a-z]+|\d+", s)


def singular(t):
    return t[:-1] if len(t) > 3 and t.endswith("s") and not t.endswith("ss") else t


def label(t):
    return t[:1].upper() + t[1:]


def analyse(data):
    keys = list(data)
    parts = {k: split_key(k) for k in keys}
    by_prefix = defaultdict(list)
    for k in keys:
        by_prefix[parts[k][0]].append(k)

    # screen prefix = the prefix whose key names other keys most often extend
    rests = {k: parts[k][1] for k in keys}
    score = Counter()
    for p, ks in by_prefix.items():
        mine = {rests[k] for k in ks if rests[k]}
        score[p] = sum(1 for k in keys if parts[k][0] != p and any(rests[k].startswith(r) for r in mine))
    screen_prefix = score.most_common(1)[0][0] if score else ""

    # features: first word, or first two words when one second word dominates
    first, second = Counter(), defaultdict(Counter)
    for k in keys:
        t = tokens(rests[k])
        if t:
            first[t[0]] += 1
            if len(t) > 1:
                second[t[0]][singular(t[1])] += 1
    two = {f: s.most_common(1)[0][0] for f, s in second.items()
           if s and s.most_common(1)[0][1] / first[f] >= 0.7}
    first_lower = {}
    for f, _ in first.most_common():
        first_lower.setdefault(f.lower(), f)

    def feature_of(k):
        t = tokens(rests[k])
        if not t:
            return "Other"
        if len(t) > 1 and (t[0] + t[1]).lower() in first_lower:
            return first_lower[(t[0] + t[1]).lower()]
        if t[0] in two and len(t) > 1 and singular(t[1]) == two[t[0]]:
            return t[0] + " " + two[t[0]]
        return first_lower.get(t[0].lower(), t[0])

    feat = {k: feature_of(k) for k in keys}
    canon = {}
    for f, _ in Counter(feat.values()).most_common():
        canon.setdefault(singular(re.sub(r"\s+", "", f).lower()), f)
    feat = {k: canon[singular(re.sub(r"\s+", "", f).lower())] for k, f in feat.items()}
    feat = {k: " ".join(tokens(f.replace(" ", ""))) if f != "Other" else f for k, f in feat.items()}

    # reading order: feature (first seen) -> screen key (file order) -> keys whose name continues it
    screens = [k for k in keys if parts[k][0] == screen_prefix]
    owner = {}
    for k in keys:
        if parts[k][0] == screen_prefix:
            continue
        c = [s for s in screens if rests[s] and rests[k].startswith(rests[s])]
        if c:
            owner[k] = max(c, key=lambda s: len(rests[s]))
    # keys in a feature that has no screen at all: join the nearest screen above them in the file
    feats_with_screens = {feat[s] for s in screens}
    pos = {k: i for i, k in enumerate(keys)}
    for k in keys:
        if k in owner or parts[k][0] == screen_prefix or feat[k] in feats_with_screens:
            continue
        for j in range(pos[k] - 1, max(-1, pos[k] - 21), -1):
            if parts[keys[j]][0] == screen_prefix:
                owner[k] = keys[j]
                feat[k] = feat[keys[j]]
                break
    children = defaultdict(list)
    for k, s in owner.items():
        children[s].append(k)
    order, seen = [], set()
    feat_order = []
    for k in keys:
        if feat[k] not in feat_order:
            feat_order.append(feat[k])
    for f in feat_order:
        for k in keys:
            if feat[k] != f or k in seen:
                continue
            if parts[k][0] == screen_prefix:
                order.append((k, 0))
                seen.add(k)
                for c in children[k]:
                    if c not in seen:
                        order.append((c, 1))
                        seen.add(c)
        for k in keys:
            if feat[k] == f and k not in seen:
                order.append((k, 0))
                seen.add(k)

    trackers = []
    for k in keys:
        v = data[k]
        if isinstance(v, dict):
            for t in v:
                if t not in trackers:
                    trackers.append(t)
    return {"keys": keys, "parts": parts, "screen_prefix": screen_prefix, "feat": feat,
            "order": order, "trackers": trackers, "feat_order": feat_order}


# ---- sheets -------------------------------------------------------------------------------------
def how_to(ws, a, blocks, rows):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Tracking documentation", "title"),
        ("", None),
        ("Documenting a flow", "h"),
        ("1. Right-click the FLOW TEMPLATE tab → Move or Copy → tick 'Create a copy' → OK. Rename the copy, "
         "e.g. 'Transfer Flow'.", None),
        ("2. Fill the flow name, owner and date at the top of the sheet.", None),
        (f"3. Each screen block has a title row (screen key + screen name) and {rows} event rows next to one "
         "merged IMAGE cell.", None),
        ("4. Screenshot: select the IMAGE cell → Insert → Pictures → Place in Cell (Excel 365). The picture "
         "stays inside the cell when rows are added or sorted.", None),
        ("5. EVENT TYPE has a drop-down with every tracker in the JSON. Write one event per row.", None),
        ("6. Several parameters or values: put each on its own line in the same cell "
         "(Windows: Alt+Enter · Mac: Option+Return), so they line up with each other.", None),
        ("7. An event fired several times with different values: list all distinct values in VALUE, one per line.", None),
        ("8. Need more rows in a screen? Insert rows inside the block (above its last row); the IMAGE cell grows "
         "with them. Need more blocks? Copy a whole empty block (title row + rows) and paste it below.", None),
        ("", None),
        ("Tracking progress", "h"),
        (f"9. When an event is documented in any flow sheet, open SUMMARY and pick {DONE_MARK} in the Done column "
         "of its row. Optionally write the flow name in 'Documented in'.", None),
        ("10. Use the filter buttons in the SUMMARY header: Done = (Blanks) gives the to-do list; filter by "
         "Feature or Type to split the work.", None),
        ("11. PROGRESS updates by itself: screens and events done, per feature.", None),
        ("", None),
        ("About SUMMARY", "h"),
        (f"• One row per key in the JSON ({len(a['keys'])} keys). Type 'Screen' = keys starting with "
         f"'{a['screen_prefix']}'; every other key is an Event, labelled with its own prefix.", None),
        ("• Feature is a guess from the first word(s) of the key name. Keys whose name continues a screen key "
         "are listed right under it, indented.", None),
        ("• Tracker columns show each tracker's event name; Parameters lists the field names the trackers receive.", None),
        ("• SUMMARY and PROGRESS only use formulas inside this workbook, so both sheets can be copied into "
         "another workbook (Move or Copy → To book).", None),
    ]
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(row=i, column=2, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            c.font = font(18, True)
        elif kind == "h":
            c.font = font(12, True)
        else:
            c.font = font(10, color="3B4357")


def flow_template(ws, trackers, blocks, rows):
    ws.sheet_view.showGridLines = False
    widths = [26, 18, 38, 30, 30, 36]
    heads = ["IMAGE", "EVENT TYPE", "EVENT NAME", "EVENT PARAMETERS", "VALUE", "NOTES"]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws["A1"] = "FLOW NAME"
    ws["A1"].font = font(18, True)
    ws["A2"] = "Owner:"
    ws["C2"] = "Last updated:"
    ws["E2"] = "Rename this sheet to the flow name"
    for ref in ("A2", "C2"):
        ws[ref].font = font(10, True, MUTED)
        ws[ref].alignment = Alignment(horizontal="right")
    ws["B2"].fill = fill("FFF8DB")
    ws["D2"].fill = fill("FFF8DB")
    ws["E2"].font = font(9, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 28

    hr = 4
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = font(11, True, "FFFFFF")
        cell.fill = fill(HEAD_BG)
        cell.alignment = CENTER
        cell.border = BOX
    ws.row_dimensions[hr].height = 26
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    dv = DataValidation(type="list", formula1='"' + ",".join(label(t) for t in trackers) + '"',
                        allow_blank=True, showErrorMessage=False)
    dv.prompt, dv.promptTitle, dv.showInputMessage = "Pick a tracker or type one", "EVENT TYPE", True
    ws.add_data_validation(dv)

    r = hr + 1
    for b in range(1, blocks + 1):
        # title row: number · screen key · screen name
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill(BLOCK_BG)
            cell.border = Border(top=Side(style="medium", color=HEAD_BG), bottom=THIN,
                                 left=THIN if c == 1 else None, right=THIN if c == 6 else None)
        ws.cell(row=r, column=1, value=f"#{b}  Screen key ›").font = font(10, True, MUTED)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(row=r, column=2).font = font(11, True, BLOCK_INK)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="center")
        for c in (2, 3, 5, 6):
            ws.cell(row=r, column=c).fill = fill("FFFFFF")
        ws.cell(row=r, column=4, value="Screen name ›").font = font(10, True, MUTED)
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        ws.cell(row=r, column=5).font = font(11, True, BLOCK_INK)
        ws.cell(row=r, column=5).alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 24
        r += 1
        # event rows
        top = r
        for i in range(rows):
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                cell.border = BOX
                cell.font = font(10, c == 2)
                cell.alignment = CENTER if c == 1 else (MID if c == 2 else TOP)
            ws.row_dimensions[r].height = 30
            r += 1
        ws.merge_cells(start_row=top, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(row=top, column=1, value="paste screenshot\n(Insert › Pictures › Place in Cell)").font = \
            font(8, italic=True, color="A0A7B4")
        dv.add(f"B{top}:B{r - 1}")
        r += 1  # spacer row between blocks
        ws.row_dimensions[r - 1].height = 10
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def summary(ws, data, a):
    ws.sheet_view.showGridLines = False
    trackers = a["trackers"]
    cols = ["Done", "Feature", "Key", "Type"] + [label(t) for t in trackers] + ["Parameters", "Documented in", "Notes"]
    widths = [7, 22, 46, 13] + [34] * len(trackers) + [30, 20, 30]
    hr = 4
    ws["A1"] = "SUMMARY · every key in the JSON"
    ws["A1"].font = font(16, True)
    ws["A2"] = (f"Pick {DONE_MARK} in Done when the key is documented in a flow sheet. Filter Done = (Blanks) for the "
                "to-do list. Keys whose name continues a screen key are indented under it.")
    ws["A2"].font = font(9, color=MUTED)
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        ws.cell(row=hr, column=c, value=h)
        ws.column_dimensions[get_column_letter(c)].width = w
    sp = a["screen_prefix"]
    r = hr + 1
    for k, level in a["order"]:
        p = a["parts"][k][0]
        v = data[k] if isinstance(data[k], dict) else {}
        names = []
        fields = []
        for t in trackers:
            sec = v.get(t)
            if isinstance(sec, dict):
                names.append(sec.get("name", ""))
                pm = sec.get("parameterMap")
                if isinstance(pm, dict):
                    for fv in pm.values():
                        if str(fv) not in fields:
                            fields.append(str(fv))
            else:
                names.append("")
        is_screen = p == sp
        row = [None, a["feat"][k], k, "Screen" if is_screen else f"Event · {p}" if p else "Event"] + names + \
              [", ".join(fields), None, None]
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = font(10, is_screen and c in (3, 4))
            cell.alignment = Alignment(vertical="top", wrap_text=c > 4,
                                       indent=2 if (c == 3 and level == 1) else 0)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        r += 1
    last = max(r - 1, hr + 1)
    lc = get_column_letter(len(cols))
    t = Table(displayName="Summary", ref=f"A{hr}:{lc}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=False)
    ws.add_table(t)
    for c in range(1, len(cols) + 1):
        ws.cell(row=hr, column=c).font = font(10, True)
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)
    dv = DataValidation(type="list", formula1=f'"{DONE_MARK}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A{hr + 1}:A{last}")
    ws.conditional_formatting.add(f"A{hr + 1}:{lc}{last}", FormulaRule(
        formula=[f'$A{hr + 1}="{DONE_MARK}"'], fill=fill("E3F5EB"), font=Font(name=F, color="1E7F4F")))
    ws.conditional_formatting.add(f"A{hr + 1}:{lc}{last}", FormulaRule(
        formula=[f'$D{hr + 1}="Screen"'], fill=fill("EEF2F8")))
    return hr, last


def progress(ws, a, hr, last):
    ws.sheet_view.showGridLines = False
    for c, w in enumerate([26, 15, 12, 14, 12, 10, 10, 16], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws["A1"] = "PROGRESS"
    ws["A1"].font = font(16, True)
    ws["A2"] = "Counts ticks in SUMMARY. Updates automatically."
    ws["A2"].font = font(9, color=MUTED)
    rng = lambda col: f"SUMMARY!${col}${hr + 1}:${col}${last}"
    A, B, D = rng("A"), rng("B"), rng("D")
    rows = [("Screens", f'=COUNTIFS({D},"Screen",{A},"{DONE_MARK}")', f'=COUNTIF({D},"Screen")'),
            ("Events", f'=COUNTIFS({D},"<>Screen",{A},"{DONE_MARK}")', f'=COUNTIF({D},"<>Screen")'),
            ("All keys", f'=COUNTIF({A},"{DONE_MARK}")', f"=ROWS({A})")]
    for c, h in enumerate(["", "Done", "Total", "Progress"], 1):
        ws.cell(row=4, column=c, value=h).font = font(10, True, MUTED)
    for i, (lbl, done, tot) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=lbl).font = font(12, True)
        ws.cell(row=i, column=2, value=done).font = font(12, True)
        ws.cell(row=i, column=3, value=tot).font = font(12)
        c = ws.cell(row=i, column=4, value=f"=IFERROR(B{i}/C{i},0)")
        c.number_format = "0%"
        c.font = font(12, True)
    hdr = ["Feature", "Screens done", "Screens", "Events done", "Events", "Done", "Total", "Progress"]
    r0 = 10
    for c, h in enumerate(hdr, 1):
        ws.cell(row=r0, column=c, value=h)
    for i, f in enumerate(a["feat_order"], r0 + 1):
        ws.cell(row=i, column=1, value=f)
        ws.cell(row=i, column=2, value=f'=COUNTIFS({B},$A{i},{D},"Screen",{A},"{DONE_MARK}")')
        ws.cell(row=i, column=3, value=f'=COUNTIFS({B},$A{i},{D},"Screen")')
        ws.cell(row=i, column=4, value=f'=COUNTIFS({B},$A{i},{D},"<>Screen",{A},"{DONE_MARK}")')
        ws.cell(row=i, column=5, value=f'=COUNTIFS({B},$A{i},{D},"<>Screen")')
        ws.cell(row=i, column=6, value=f"=B{i}+D{i}")
        ws.cell(row=i, column=7, value=f"=C{i}+E{i}")
        ws.cell(row=i, column=8, value=f"=IFERROR(F{i}/G{i},0)").number_format = "0%"
        for c in range(1, 9):
            ws.cell(row=i, column=c).font = font(10)
    end = r0 + len(a["feat_order"])
    t = Table(displayName="Progress", ref=f"A{r0}:H{end}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
    ws.add_table(t)
    for c in range(1, 9):
        ws.cell(row=r0, column=c).font = font(10, True)
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add(f"H{r0 + 1}:H{end}", DataBarRule(start_type="num", start_value=0,
                                  end_type="num", end_value=1, color="5B8DEF"))
    ws.conditional_formatting.add("D5:D7", DataBarRule(start_type="num", start_value=0, end_type="num",
                                  end_value=1, color="5B8DEF"))
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json")
    ap.add_argument("--out", default="tracking_workbook.xlsx")
    ap.add_argument("--blocks", type=int, default=20, help="screen blocks in FLOW TEMPLATE (default 20)")
    ap.add_argument("--rows", type=int, default=8, help="event rows per screen block (default 8)")
    args = ap.parse_args()
    with open(args.json, encoding="utf-8") as fh:
        data = json.load(fh, object_pairs_hook=OrderedDict)
    if not isinstance(data, dict):
        sys.exit("top level of the JSON is not an object")
    a = analyse(data)
    wb = Workbook()
    how_to(wb.active, a, args.blocks, args.rows)
    wb.active.title = "HOW TO USE"
    flow_template(wb.create_sheet("FLOW TEMPLATE"), a["trackers"], args.blocks, args.rows)
    hr, last = summary(wb.create_sheet("SUMMARY"), data, a)
    progress(wb.create_sheet("PROGRESS"), a, hr, last)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.out)
    n_screens = sum(1 for k in a["keys"] if a["parts"][k][0] == a["screen_prefix"])
    print(f"wrote {args.out}: {len(a['keys'])} keys ({n_screens} screens, {len(a['keys']) - n_screens} events), "
          f"{len(a['feat_order'])} features, {len(a['trackers'])} trackers, {args.blocks} screen blocks in the template")


if __name__ == "__main__":
    main()
