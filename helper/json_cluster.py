#!/usr/bin/env python3
"""Clusters the keys of a tracking-spec JSON file into features and screens, prints a short
preview, and writes an Excel file to check the result.

  pip install openpyxl
  python3 json_cluster.py <file.json>                       # writes json_clusters.xlsx
  python3 json_cluster.py <file.json> --aliases a.json      # optional feature renames/merges
  python3 json_cluster.py <file.json> --out other.xlsx

How keys are clustered (everything is detected from the file, nothing is hard-coded):
  prefix   = the leading lowercase word of a key ("abcDefGhi" -> "abc")
  screen   = keys of the "screen prefix": the prefix whose key names other keys most often
             extend (detected automatically, override with --screen-prefix)
  feature  = first word after the prefix, or the first two words when one second word
             dominates (e.g. "Credit"+"Card" in most keys -> "Credit Card")
  placement of every other key, in order:
    1. name match   : its name continues a screen key's name (longest match wins)
    2. file order   : nearest screen key above it in the file, same feature, within --window keys
    3. unassigned

--aliases is a JSON object mapping a detected feature name to the name you want, e.g.
  {"Feature Names": "Feature Name"}  (merge a plural into the singular)
"""
import argparse
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    sys.exit("needs openpyxl:  pip install openpyxl")

DOMINANT_SECOND = 0.7      # share of a first-word group that must share the second word
W = 100


def split_key(k):
    m = re.match(r"^([a-z]+)(.*)$", k)
    return (m.group(1), m.group(2)) if m else ("", k)


def tokens(s):
    return re.findall(r"[A-Z]+(?=[A-Z][a-z]|\d|$)|[A-Z]?[a-z]+|\d+", s)


def humanize(s):
    return " ".join(tokens(s)) or s


def singular(t):
    return t[:-1] if len(t) > 3 and t.endswith("s") and not t.endswith("ss") else t


def head(title):
    print("\n" + ("── " + title + " ").ljust(W, "─"))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json")
    ap.add_argument("--out", default="json_clusters.xlsx")
    ap.add_argument("--aliases", help="JSON object: detected feature name -> wanted name")
    ap.add_argument("--screen-prefix", help="override the detected screen prefix")
    ap.add_argument("--window", type=int, default=20, help="max keys between a key and its screen for file-order placement")
    args = ap.parse_args()

    with open(args.json, encoding="utf-8") as fh:
        data = json.load(fh, object_pairs_hook=OrderedDict)
    if not isinstance(data, dict):
        sys.exit("top level is not an object")
    aliases = {}
    if args.aliases:
        with open(args.aliases, encoding="utf-8") as fh:
            aliases = json.load(fh)

    keys = list(data)
    pos = {k: i for i, k in enumerate(keys)}
    parts = {k: split_key(k) for k in keys}
    by_prefix = defaultdict(list)
    for k in keys:
        by_prefix[parts[k][0]].append(k)

    # ---- screen prefix: the one whose names others extend most ----------------------------------
    if args.screen_prefix:
        screen_prefix = args.screen_prefix
    else:
        base_score = Counter()
        rests = {k: parts[k][1] for k in keys}
        for p, ks in by_prefix.items():
            mine = {rests[k] for k in ks if rests[k]}
            base_score[p] = sum(1 for k in keys if parts[k][0] != p and any(rests[k].startswith(r) for r in mine))
        screen_prefix = base_score.most_common(1)[0][0] if base_score else ""
    screens = [k for k in keys if parts[k][0] == screen_prefix]

    # ---- features -----------------------------------------------------------------------------------
    first = Counter()
    second_by_first = defaultdict(Counter)
    for k in keys:
        t = tokens(parts[k][1])
        if t:
            first[t[0]] += 1
            if len(t) > 1:
                second_by_first[t[0]][singular(t[1])] += 1
    two_word = {}
    for f, n in first.items():
        if second_by_first[f]:
            s, c = second_by_first[f].most_common(1)[0]
            if c / n >= DOMINANT_SECOND:
                two_word[f] = s

    # same word written differently ("AbcDef" vs "Abcdef"): compare without case and word breaks
    first_lower = {}
    for f, n in first.most_common():
        first_lower.setdefault(f.lower(), f)

    def feature_of(k):
        t = tokens(parts[k][1])
        if not t:
            return "(no name)"
        name = t[0]
        if len(t) > 1 and (t[0] + t[1]).lower() in first_lower:
            name = first_lower[(t[0] + t[1]).lower()]
        elif t[0] in two_word and len(t) > 1 and singular(t[1]) == two_word[t[0]]:
            name = t[0] + " " + two_word[t[0]]
        else:
            name = first_lower.get(t[0].lower(), t[0])
        return aliases.get(name, name)

    feat = {k: feature_of(k) for k in keys}
    # merge features that differ only in case / spacing / a plural "s"
    canon = {}
    for f, n in Counter(feat.values()).most_common():
        canon.setdefault(singular(re.sub(r"\s+", "", f).lower()), f)
    feat = {k: canon[singular(re.sub(r"\s+", "", f).lower())] for k, f in feat.items()}

    # ---- placement ------------------------------------------------------------------------------------
    screen_rest = {s: parts[s][1] for s in screens}
    placed = {}     # key -> (screen key or None, how)
    for s in screens:
        placed[s] = (s, "screen")
    for k in keys:
        if k in placed:
            continue
        rest = parts[k][1]
        cands = [s for s in screens if screen_rest[s] and rest.startswith(screen_rest[s])]
        if cands:
            placed[k] = (max(cands, key=lambda s: len(screen_rest[s])), "name match")
    for k in keys:
        if k in placed:
            continue
        i = pos[k]
        found = None
        for j in range(i - 1, max(-1, i - 1 - args.window), -1):
            c = keys[j]
            if parts[c][0] == screen_prefix and feat[c] == feat[k]:
                found = c
                break
        placed[k] = (found, "file order") if found else (None, "unassigned")

    how_count = Counter(how for _, how in placed.values())
    others = len(keys) - len(screens)

    # ---- terminal preview -----------------------------------------------------------------------------
    head("CLUSTERING PREVIEW")
    print(f"  {len(keys)} keys · screen prefix: '{screen_prefix}' ({len(screens)} screens) · {others} other keys")
    for how in ["name match", "file order", "unassigned"]:
        n = how_count.get(how, 0)
        print(f"    {how:<12} {n:>5}  ({n / max(1, others) * 100:.0f}% of other keys)")
    feat_count = Counter(feat.values())
    feat_screens = Counter(feat[s] for s in screens)
    feat_un = Counter(feat[k] for k, (_, h) in placed.items() if h == "unassigned")
    head(f"FEATURES ({len(feat_count)})")
    items = [f"{f} {n}k/{feat_screens.get(f, 0)}s" + (f"/{feat_un[f]}u" if feat_un.get(f) else "")
             for f, n in feat_count.most_common()]
    line = " "
    for it in items[:60]:
        if len(line) + len(it) + 3 > W:
            print(line)
            line = " "
        line += " " + it + " ·"
    print(line.rstrip(" ·"))
    print("  (k = keys, s = screens, u = unassigned)" + (f" · {len(items) - 60} more in Excel" if len(items) > 60 else ""))

    head("SAMPLE CLUSTERS")
    members = defaultdict(list)
    for k, (s, how) in placed.items():
        if s and k != s:
            members[s].append((k, how))
    sample = sorted(members, key=lambda s: -len(members[s]))[:3]
    for s in sample:
        print(f"  [{feat[s]}] {s}")
        for k, how in sorted(members[s], key=lambda x: pos[x[0]])[:6]:
            print(f"      {k:<55} {how}")
        if len(members[s]) > 6:
            print(f"      … {len(members[s]) - 6} more")

    head("30 CONSECUTIVE KEYS FROM THE MIDDLE OF THE FILE (prefix · key)")
    mid = max(0, len(keys) // 2 - 15)
    for k in keys[mid:mid + 30]:
        print(f"  {pos[k]:>4}  {parts[k][0]:<10} {k}")

    # ---- Excel ------------------------------------------------------------------------------------------
    trackers = []
    for k in keys:
        v = data[k]
        if isinstance(v, dict):
            for t in v:
                if t not in trackers:
                    trackers.append(t)
    feat_order = {}
    for k in keys:
        feat_order.setdefault(feat[k], len(feat_order))

    def row_sort(k):
        s, how = placed[k]
        anchor = pos[s] if s else 10 ** 6 + pos[k]
        return (feat_order[feat[k]], 0 if s else 1, anchor, 0 if how == "screen" else 1, pos[k])

    ordered = sorted(keys, key=row_sort)
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    wb = Workbook()

    # OVERVIEW
    ov = wb.active
    ov.title = "OVERVIEW"
    ov["A1"] = "Key clustering preview"
    ov["A1"].font = Font(name="Arial", size=16, bold=True)
    ov["A2"] = f"source: {args.json.split('/')[-1]} · {len(keys)} keys · screen prefix '{screen_prefix}'"
    ov["A2"].font = Font(name="Arial", size=9, color="5B6478")
    rows = [("Screens", len(screens)), ("Other keys", others)] + \
           [(f"  placed by {h}", how_count.get(h, 0)) for h in ["name match", "file order"]] + \
           [("  unassigned", how_count.get("unassigned", 0))]
    for i, (a, b) in enumerate(rows, 4):
        ov.cell(row=i, column=1, value=a).font = font
        ov.cell(row=i, column=2, value=b).font = bold
    r0 = 4 + len(rows) + 1
    hdr = ["Feature", "Keys", "Screens", "Name match", "File order", "Unassigned"]
    for c, h in enumerate(hdr, 1):
        ov.cell(row=r0, column=c, value=h).font = hdr_font
    fstats = defaultdict(Counter)
    for k in keys:
        fstats[feat[k]]["keys"] += 1
        fstats[feat[k]][placed[k][1]] += 1
    for i, f in enumerate(sorted(fstats, key=lambda f: feat_order[f]), r0 + 1):
        st = fstats[f]
        vals = [f, st["keys"], st["screen"], st["name match"], st["file order"], st["unassigned"]]
        for c, v in enumerate(vals, 1):
            ov.cell(row=i, column=c, value=v).font = font
    last = r0 + len(fstats)
    t = Table(displayName="Features", ref=f"A{r0}:F{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ov.add_table(t)
    for c, w in enumerate([28, 10, 10, 12, 12, 12], 1):
        ov.column_dimensions[get_column_letter(c)].width = w
    ov.cell(row=last + 2, column=1, value="Sheets: CLUSTERS = one row per key per tracker, grouped by feature and "
            "screen · UNASSIGNED = keys with no screen · FILE ORDER = every key in file order").font = \
        Font(name="Arial", size=9, color="5B6478")

    def write_keys_sheet(ws, name, ks):
        cols = ["Feature", "Screen", "Screen key", "Key", "Prefix", "Placed by", "Tracker", "Event name", "Parameters", "File #"]
        widths = [22, 34, 40, 48, 10, 13, 16, 44, 40, 8]
        for c, (h, w) in enumerate(zip(cols, widths), 1):
            ws.cell(row=1, column=c, value=h).font = hdr_font
            ws.column_dimensions[get_column_letter(c)].width = w
        r = 2
        for k in ks:
            s, how = placed[k]
            v = data[k] if isinstance(data[k], dict) else {}
            secs = [(tr, sec) for tr, sec in v.items() if isinstance(sec, dict)] or [("", {})]
            for tr, sec in secs:
                pm = sec.get("parameterMap") if isinstance(sec.get("parameterMap"), dict) else {}
                vals = [feat[k], humanize(parts[s][1]) if s else "(unassigned)", s or "", k, parts[k][0], how,
                        tr, sec.get("name", ""), ", ".join(str(x) for x in pm.values()), pos[k] + 1]
                for c, val in enumerate(vals, 1):
                    ws.cell(row=r, column=c, value=val).font = bold if how == "screen" else font
                r += 1
        last_r = max(r - 1, 2)
        tb = Table(displayName=name, ref=f"A1:{get_column_letter(len(cols))}{last_r}")
        tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(tb)
        ws.freeze_panes = "E2"
        ws.conditional_formatting.add(f"A2:J{last_r}", FormulaRule(formula=['$F2="screen"'],
                                      fill=PatternFill("solid", fgColor="DCE3EE")))
        ws.conditional_formatting.add(f"A2:J{last_r}", FormulaRule(formula=['$F2="unassigned"'],
                                      fill=PatternFill("solid", fgColor="FDE2E1")))
        ws.conditional_formatting.add(f"A2:J{last_r}", FormulaRule(formula=['$F2="file order"'],
                                      fill=PatternFill("solid", fgColor="FEF6E0")))

    write_keys_sheet(wb.create_sheet("CLUSTERS"), "Clusters", ordered)
    write_keys_sheet(wb.create_sheet("UNASSIGNED"), "Unassigned",
                     [k for k in ordered if placed[k][1] == "unassigned"])

    fo = wb.create_sheet("FILE ORDER")
    cols = ["File #", "Key", "Prefix", "Feature", "Screen key", "Placed by"]
    for c, (h, w) in enumerate(zip(cols, [8, 50, 10, 22, 44, 13]), 1):
        fo.cell(row=1, column=c, value=h).font = hdr_font
        fo.column_dimensions[get_column_letter(c)].width = w
    for r, k in enumerate(keys, 2):
        s, how = placed[k]
        for c, val in enumerate([pos[k] + 1, k, parts[k][0], feat[k], s or "", how], 1):
            fo.cell(row=r, column=c, value=val).font = bold if how == "screen" else font
    tb = Table(displayName="FileOrder", ref=f"A1:F{len(keys) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    fo.add_table(tb)
    fo.freeze_panes = "C2"

    wb.save(args.out)
    print(f"\n  wrote {args.out}  (OVERVIEW · CLUSTERS · UNASSIGNED · FILE ORDER)\n")


if __name__ == "__main__":
    main()
